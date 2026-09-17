import streamlit as st
import pypdf
import json
import io
import google.generativeai as genai
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from typing import List, Dict

# --- UI CONFIGURATION ---
st.set_page_config(page_title="AI PDF to PPT Converter", page_icon="📊", layout="wide")

# Custom CSS for a professional look
st.markdown("""
    <style>
    .main { background-color: #0f172a; color: white; }
    .stButton>button { 
        width: 100%; 
        background-color: #14b8a6; 
        color: white; 
        border: none; 
        padding: 0.75rem;
        font-weight: bold;
        border-radius: 0.5rem;
    }
    .stButton>button:hover { background-color: #0d9488; border: none; color: white; }
    div[data-testid="stExpander"] { border: 1px solid #334155; border-radius: 0.5rem; }
    </style>
    """, unsafe_allow_html=True)

# --- CORE LOGIC ---

def extract_text_from_pdf(uploaded_file) -> str:
    """Extracts text from PDF with robust error handling."""
    try:
        reader = pypdf.PdfReader(uploaded_file)
        full_text = []
        for page in reader.pages:
            content = page.extract_text()
            if content:
                full_text.append(content)
        
        if not full_text:
            return None
        return "\n".join(full_text)
    except Exception as e:
        st.error(f"Error reading PDF: {e}")
        return None

def call_gemini_api(api_key: str, text: str, style: str, custom_instr: str) -> List[Dict]:
    """Calls Gemini API and ensures JSON output."""
    try:
        genai.configure(api_key=api_key)
        # Using 1.5-flash for speed and reliability
        model = genai.GenerativeModel('gemini-1.5-flash')
        
        system_instruction = (
            "You are a professional chemistry professor and slide designer. "
            "Summarize the provided text into a 6-10 slide presentation. "
            "Output MUST be a valid JSON array of objects. "
            "Format: [{\"title\": \"slide title\", \"content\": [\"bullet 1\", \"bullet 2\"], \"notes\": \"speaker notes\"}]"
        )
        
        prompt = f"Style: {style}. Extra Instructions: {custom_instr}. Text: {text[:30000]}"
        
        response = model.generate_content(
            f"{system_instruction}\n\n{prompt}",
            generation_config={"response_mime_type": "application/json"}
        )
        
        return json.loads(response.text)
    except Exception as e:
        st.error(f"AI Error: {e}")
        return None

def create_pptx(slides_data: List[Dict], style: str):
    """Builds PPTX file in memory."""
    prs = Presentation()
    
    for slide_data in slides_data:
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        
        # Title
        title = slide.shapes.title
        title.text = slide_data.get('title', 'Slide')
        
        # Content
        content_placeholder = slide.placeholders[1]
        tf = content_placeholder.text_frame
        tf.word_wrap = True
        
        for bullet in slide_data.get('content', []):
            p = tf.add_paragraph()
            p.text = bullet
            p.level = 0
            
        # Notes
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = slide_data.get('notes', '')

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer

# --- MAIN APP ---

def main():
    st.title("🧪 AI Chemistry PDF to PPT Converter")
    st.write("Convert textbook chapters into AI-structured presentations instantly.")

    with st.sidebar:
        st.header("1. Setup")
        api_key = st.text_input("Enter Gemini API Key", type="password")
        st.markdown("[How to get a key?](https://aistudio.google.com/app/apikey)")
        
        st.header("2. Design")
        style = st.selectbox("Presentation Style", ["Academic", "Gen Z", "Formal", "Infographic"])
        custom_instr = st.text_area("Custom Instructions", placeholder="e.g. Focus on the Chemical Kinetics formulas...")

    st.header("3. Upload")
    uploaded_file = st.file_uploader("Choose a PDF file", type="pdf")

    if uploaded_file:
        st.info(f"📄 Selected: {uploaded_file.name}")
        
        if st.button("Generate Presentation"):
            if not api_key:
                st.warning("Please add your Gemini API key in the sidebar.")
                return

            with st.status("Processing...", expanded=True) as status:
                # Step 1: Text
                status.update(label="Extracting Text from PDF...")
                text = extract_text_from_pdf(uploaded_file)
                if not text:
                    status.update(label="Failed to extract text.", state="error")
                    return

                # Step 2: AI
                status.update(label="AI Summarizing Content...")
                slides_json = call_gemini_api(api_key, text, style, custom_instr)
                if not slides_json:
                    status.update(label="AI Summary failed.", state="error")
                    return

                # Step 3: PPTX
                status.update(label="Creating PowerPoint File...")
                ppt_buffer = create_pptx(slides_json, style)
                
                status.update(label="Done!", state="complete", expanded=False)

            st.success("Your presentation is ready!")
            st.download_button(
                label="Download PowerPoint (.pptx)",
                data=ppt_buffer,
                file_name=f"{uploaded_file.name.replace('.pdf', '')}_Summary.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
            
            with st.expander("View Slide Outline"):
                st.json(slides_json)

if __name__ == "__main__":
    main()
